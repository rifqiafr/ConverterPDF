import os
import sys
import docx
import json
import subprocess

TEST_DIR = os.path.join(os.path.dirname(__file__), 'test_files')
os.makedirs(TEST_DIR, exist_ok=True)

ENGINE = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'engine', 'convert.py'))

def run_engine(action, input_files, output_file, extra_args=None):
    cmd = [sys.executable, ENGINE, '--action', action, '--input', input_files, '--output', output_file]
    if extra_args:
        cmd.extend(extra_args)
    res = subprocess.run(cmd, capture_output=True, text=True)
    try:
        lines = [line.strip() for line in res.stdout.strip().splitlines() if line.strip()]
        for line in reversed(lines):
            if line.startswith('{') and line.endswith('}'):
                return json.loads(line)
        return {"success": False, "raw": res.stdout, "stderr": res.stderr}
    except Exception:
        return {"success": False, "raw": res.stdout, "stderr": res.stderr}

def main():
    print("=== MULAI PENGUJIAN OTOMATIS MODUL PDF CONVERTER ===")
    
    # 1. Buat file sampel DOCX
    docx_path = os.path.join(TEST_DIR, 'sample.docx')
    doc = docx.Document()
    doc.add_heading('Laporan Uji Coba Konversi Dokumen', 0)
    p = doc.add_paragraph('Ini adalah paragraf contoh dengan teks ')
    p.add_run('tebal').bold = True
    p.add_run(' dan teks ')
    p.add_run('miring.').italic = True
    
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'No'
    hdr_cells[1].text = 'Nama Modul'
    hdr_cells[2].text = 'Status'
    
    row_cells = table.add_row().cells
    row_cells[0].text = '1'
    row_cells[1].text = 'Word to PDF'
    row_cells[2].text = 'Siap Uji'
    doc.save(docx_path)
    print(f"[OK] File sampel Word dibuat: {docx_path}")

    # 2. Uji Word ke PDF
    pdf_path = os.path.join(TEST_DIR, 'sample_from_word.pdf')
    res1 = run_engine('word_to_pdf', docx_path, pdf_path)
    print("1. Word to PDF:", "[BERHASIL]" if res1.get('success') else f"[GAGAL: {res1}]")

    # 3. Uji PDF ke Word
    docx_back = os.path.join(TEST_DIR, 'sample_from_pdf.docx')
    res2 = run_engine('pdf_to_word', pdf_path, docx_back)
    print("2. PDF to Word:", "[BERHASIL]" if res2.get('success') else f"[GAGAL: {res2}]")

    # 4. Uji Merge PDF (gabung 2 file)
    merged_pdf = os.path.join(TEST_DIR, 'merged_output.pdf')
    res3 = run_engine('merge_pdf', f"{pdf_path},{pdf_path}", merged_pdf)
    print("3. Merge PDF:", "[BERHASIL]" if res3.get('success') else f"[GAGAL: {res3}]")

    # 5. Uji Split PDF
    split_zip = os.path.join(TEST_DIR, 'split_output.zip')
    res4 = run_engine('split_pdf', merged_pdf, split_zip)
    print("4. Split PDF:", "[BERHASIL]" if res4.get('success') else f"[GAGAL: {res4}]")

    # 6. Uji Compress PDF
    compressed_pdf = os.path.join(TEST_DIR, 'compressed_output.pdf')
    res5 = run_engine('compress_pdf', merged_pdf, compressed_pdf, extra_args=['--quality', 'medium'])
    print("5. Compress PDF:", "[BERHASIL]" if res5.get('success') else f"[GAGAL: {res5}]")

    # 7. Uji PDF to JPG
    jpg_zip = os.path.join(TEST_DIR, 'pdf_to_images.zip')
    res6 = run_engine('pdf_to_jpg', pdf_path, jpg_zip)
    print("6. PDF to Images (JPG):", "[BERHASIL]" if res6.get('success') else f"[GAGAL: {res6}]")

    # 8. Uji PDF to Excel
    excel_path = os.path.join(TEST_DIR, 'extracted_table.xlsx')
    res7 = run_engine('pdf_to_excel', pdf_path, excel_path)
    print("7. PDF to Excel:", "[BERHASIL]" if res7.get('success') else f"[GAGAL: {res7}]")

    # 9. Uji PDF to PPTX
    pptx_path = os.path.join(TEST_DIR, 'presentation.pptx')
    res8 = run_engine('pdf_to_pptx', pdf_path, pptx_path)
    print("8. PDF to PPTX:", "[BERHASIL]" if res8.get('success') else f"[GAGAL: {res8}]")

    print("\n=== SEMUA MODUL PENGUJIAN SELESAI DILAKUKAN ===")

if __name__ == '__main__':
    main()
